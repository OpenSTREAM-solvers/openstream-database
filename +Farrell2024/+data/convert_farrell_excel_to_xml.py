#!/usr/bin/env python3
"""Convert Farrell 2024 Excel film-thickness data into dataset XML."""

import sys
from pathlib import Path
from xml.etree import ElementTree as ET
import openpyxl

FIELD_MAP = {
    "Param Name": "ParamName",
    "Mass Flow [kg/m^2]": "MassFlow",
    "Mass Flux [kg/m^2*s]": "MassFlux",
    "Inlet Enthalpy [J/kg]": "InletEnthalpy",
    "Inlet Quality": "InletQuality",
    "Heat Flux [W/m^2]": "HeatFlux",
    "Combined Heater Power [W]": "Power",
    "Pressure [Pa]": "Pressure",
    "Length [m]": "Length",
    "Area [m^2]": "Area",
    "Perimeter [m]": "Perimeter",
    "WallMesh": "WallMesh",
    "WallPower": "WallPower",
    "Fluid": "Fluid",
    "Film Thickness Measured [m]": "MeasuredFilmThickness",
}

DEFAULT_FIELDS = {
    "Fluid": "R245fa",
}


def parse_value(raw):
    value = raw.strip()
    if value == "":
        return None
    try:
        if value.upper() in {"NA", "N/A", "NONE"}:
            return None
        return float(value)
    except ValueError:
        return value


def parse_array(raw):
    """Parse array values from Excel cells or strings."""
    if raw is None:
        return []
    
    # If it's a numpy array or list-like object
    if hasattr(raw, '__iter__') and not isinstance(raw, (str, bytes)):
        parsed_items = []
        for item in raw:
            if item is None:
                continue
            try:
                num = float(item)
                parsed_items.append(str(int(num)) if num.is_integer() else str(num))
            except (ValueError, TypeError):
                parsed_items.append(str(item))
        return parsed_items
    
    # Handle string representation
    value = str(raw).strip()
    if not value:
        return []
    if value.startswith("[") and value.endswith("]"):
        inner = value[1:-1].strip()
        if not inner:
            return []
        items = inner.split()
        parsed_items = []
        for item in items:
            try:
                num = float(item)
                parsed_items.append(str(int(num)) if num.is_integer() else str(num))
            except ValueError:
                parsed_items.append(item)
        return parsed_items
    return [value]


def pretty_print_xml(elem, level=0):
    indent = "    "
    i = "\n" + level * indent
    if len(elem):
        if not elem.text or not elem.text.strip():
            elem.text = i + indent
        for child in elem:
            pretty_print_xml(child, level + 1)
        if not child.tail or not child.tail.strip():
            child.tail = i
    if level and (not elem.tail or not elem.tail.strip()):
        elem.tail = i


def clean_param_name(value):
    if value is None:
        return ""
    return value.strip().strip("'\"")


def build_dataset_row(row, index):
    dataset = {
        "TestID": str(index),
    }

    for old_name, new_name in FIELD_MAP.items():
        if old_name not in row:
            continue
        raw_value = row[old_name]
        
        if new_name == "ParamName":
            dataset[new_name] = clean_param_name(raw_value)
            dataset["TestName"] = clean_param_name(raw_value)
        elif new_name in {"Perimeter", "WallMesh", "WallPower"}:
            # Parse array values from Excel
            array_values = parse_array(raw_value)
            dataset[new_name] = array_values
        else:
            parsed = parse_value(raw_value)
            dataset[new_name] = str(parsed) if parsed is not None else ""

    for key, value in DEFAULT_FIELDS.items():
        dataset.setdefault(key, value)

    return dataset


def write_xml(rows, xml_path):
    root = ET.Element("struct")
    for row in rows:
        dataset_elem = ET.SubElement(root, "dataset")
        for key, value in row.items():
            if isinstance(value, list):
                for item in value:
                    child = ET.SubElement(dataset_elem, key)
                    child.text = item
            else:
                child = ET.SubElement(dataset_elem, key)
                child.text = value

    pretty_print_xml(root)
    tree = ET.ElementTree(root)
    tree.write(str(xml_path), encoding="utf-8", xml_declaration=True)


def load_excel(excel_path):
    """Load data from an Excel file and return list of dictionaries."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    
    # Get header row
    header_row = []
    for cell in ws[1]:
        if cell.value is not None:
            header_row.append(str(cell.value).strip())
    
    if not header_row or "Param Name" not in header_row:
        raise ValueError("Could not find the Excel header row containing 'Param Name'.")
    
    # Read data rows
    dict_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Skip empty rows
        if not any(cell is not None for cell in row):
            continue
        
        # Create dictionary mapping headers to values
        row_dict = {}
        for i, header in enumerate(header_row):
            if i < len(row):
                value = row[i]
                # Convert to string, handling None values
                row_dict[header] = str(value) if value is not None else ""
        
        dict_rows.append(row_dict)
    
    return dict_rows


def main(argv=None):
    if argv is None:
        argv = sys.argv[1:]

    script_dir = Path(__file__).parent

    if len(argv) not in {0, 1, 2}:
        print("Usage: python convert_farrell_csv_to_xml.py [<input.xlsx>] [<output.xml>]")
        sys.exit(1)

    excel_path = Path(argv[0]) if len(argv) >= 1 else script_dir / "PreObsFilmThicknessComparison.xlsx"
    xml_path = Path(argv[1]) if len(argv) == 2 else script_dir.parent / "+src" / "Farrell2024.xml"

    rows = load_excel(excel_path)
    converted = [build_dataset_row(row, idx + 1) for idx, row in enumerate(rows)]
    write_xml(converted, xml_path)
    print(f"Wrote {len(converted)} dataset entries to {xml_path}")


if __name__ == "__main__":
    main()
